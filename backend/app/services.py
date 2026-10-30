from io import BytesIO
from pathlib import Path
from threading import RLock

from .models import (
    Artifact,
    ArtifactType,
    ContextItem,
    ContextPackage,
    Project,
    ProjectMemory,
    ProjectState,
    SDLCStage,
    TaskStatus,
    uid,
)
from .router import ARTIFACT_STAGE, STAGE_ARTIFACT


class ApplicationService:
    def __init__(self, runtime):
        self.runtime = runtime
        self.storage = runtime.storage
        self._locks = {}

    def lock(self, project_id):
        return self._locks.setdefault(project_id, RLock())

    def create_project(
        self, name, stage="Planning", requirement="", description="", model_policy="user_default"
    ):
        self.runtime.settings.current(model_policy)
        project = Project(project_name=name.strip(), description=description, model_policy=model_policy)
        stage = SDLCStage(stage)
        state = ProjectState(
            project_id=project.project_id,
            current_stage=stage,
            expected_next_artifact=STAGE_ARTIFACT[stage].value,
            requirement=requirement,
            started_stages=[stage],
            recommended_next_action=f"Generate {STAGE_ARTIFACT[stage].value}",
        )
        with self.storage.transaction(project.project_id):
            self.storage.create_project(project)
            self.storage.save_state(project.project_id, state)
            self.storage.save_memory(project.project_id, ProjectMemory(project_id=project.project_id))
            self.runtime.memory.refresh(state)
        return self.open_project(project.project_id)

    def list_projects(self):
        result = []
        for p in self.storage.list_projects():
            with self.storage.transaction(p.project_id):
                status = self.runtime.status.get(p.project_id)
                result.append({**p.model_dump(mode="json"), "status": status})
        return result

    def open_project(self, project_id):
        with self.lock(project_id), self.storage.transaction(project_id):
            project = self.storage.load_project(project_id)
            state = self.storage.load_state(project_id)
            if state.task_status in (
                TaskStatus.EXECUTING,
                TaskStatus.VALIDATING,
                TaskStatus.PLANNING_CONTEXT,
                TaskStatus.READY_TO_EXECUTE,
            ):
                state.task_status = TaskStatus.FAILED
                state.recommended_next_action = "Retry interrupted task using its existing task ID"
                self.storage.save_state(project_id, state)
            return {
                "project": project,
                "state": state,
                "memory": self.storage.load_memory(project_id),
                "status": self.runtime.status.get(project_id),
                "session": self.storage.load_session(project_id),
                "context": self.runtime.context_session.status(project_id),
                "assets": self.storage.list_project_assets(project_id),
            }

    def _message(self, project_id, role, text, metadata=None):
        with self.storage.transaction(project_id):
            self.runtime.context_session.append(project_id, role, text, metadata)

    def chat(
        self,
        project_id,
        message="",
        action=None,
        artifact_type=None,
        task_id=None,
        clarification_id=None,
        version=None,
        feedback="",
        stage=None,
        model_policy=None,
        progress=None,
    ):
        with self.lock(project_id):
            self.storage.load_project(project_id)
            if message:
                self._message(project_id, "user", message)
            state = self.storage.load_state(project_id)
            text = message.strip().lower()
            aliases = [
                ("BRD", ("brd", "business requirement")),
                ("FRD", ("frd", "functional requirement", "user stories")),
                ("TDD", ("tdd", "technical design")),
                ("ABAP", ("abap", "code generation")),
                ("UAT", ("uat", "acceptance test")),
                ("AUDIT_REPORT", ("audit", "code review")),
            ]
            if not action:
                if state.task_status == TaskStatus.AWAITING_INPUT and text:
                    action = "answer"
                elif not text or text in ("status", "project status", "/status"):
                    action = "status"
                elif text.startswith(("requirement:", "cr:")):
                    action = "submit_requirement"
                elif (
                    text.rstrip("!.? ")
                    in (
                        "hi",
                        "hello",
                        "hey",
                        "good morning",
                        "good afternoon",
                        "good evening",
                        "thanks",
                        "thank you",
                    )
                    or text.endswith("?")
                    or text.startswith(
                        (
                            "what ",
                            "how ",
                            "why ",
                            "who ",
                            "which ",
                            "where ",
                            "when ",
                            "explain ",
                            "describe ",
                            "tell me ",
                            "please explain ",
                            "help me understand ",
                            "can you ",
                            "could you ",
                            "would you ",
                        )
                    )
                ):
                    action = "conversation"
                elif not state.requirement and (
                    len(message) > 120
                    or "\n" in message
                    or not any(w in text for _, words in aliases for w in words)
                ):
                    action = "submit_requirement"
                elif text.startswith("retry"):
                    action, artifact_type, task_id = (
                        "generate",
                        state.expected_next_artifact,
                        state.active_task_id,
                    )
                elif any(word in text for word in ("generate", "create ", "revise", "draft ", "prepare")):
                    action = "generate"
                    artifact_type = next(
                        (kind for kind, words in aliases if any(w in text for w in words)),
                        state.expected_next_artifact,
                    )
                elif text.startswith(("approve ", "reject ")):
                    action = "approve" if text.startswith("approve") else "reject"
                    artifact_type = next(
                        (kind for kind, words in aliases if any(w in text for w in words)), None
                    )
                elif not state.requirement or text.startswith(("requirement:", "cr:")):
                    action = "submit_requirement"
                else:
                    action = "conversation"
            result = {}
            self.runtime.orchestrator.events.callback = progress
            try:
                if action == "conversation":
                    result = self.runtime.orchestrator.converse(project_id, message, model_policy, task_id)
                elif action == "generate":
                    project = self.storage.load_project(project_id)
                    result = self.runtime.orchestrator.run(
                        project_id,
                        artifact_type,
                        inputs={
                            "instructions": feedback,
                            "model_policy": model_policy or project.model_policy,
                        },
                        task_id=task_id or uid(),
                    )
                elif action == "answer":
                    pending = [c for c in state.open_clarifications if c.status == "open"]
                    if not pending:
                        raise ValueError("No clarification is awaiting an answer")
                    result = self.runtime.orchestrator.clarifications.answer(
                        self.runtime.orchestrator,
                        project_id,
                        clarification_id or pending[0].clarification_id,
                        message,
                    )
                elif action in ("approve", "reject"):
                    if not artifact_type:
                        raise ValueError("Choose an artifact to review")
                    artifact = self.storage.load_artifact(project_id, artifact_type, version)
                    artifact = self.runtime.approvals.review(
                        project_id,
                        artifact_type,
                        artifact.artifact_version,
                        action == "approve",
                        feedback or (message if action == "reject" else ""),
                    )
                    result = {"artifact": artifact}
                elif action == "submit_requirement":
                    if not message.strip():
                        raise ValueError("Requirement cannot be empty")
                    with self.storage.transaction(project_id):
                        state.requirement = message
                        self.storage.save_state(project_id, state)
                        self.runtime.memory.refresh(state)
                    result = {"message": "Requirement captured. Generate the BRD when ready."}
                elif action == "set_stage":
                    with self.storage.transaction(project_id):
                        state.current_stage = SDLCStage(stage)
                        if state.current_stage not in state.started_stages:
                            state.started_stages.append(state.current_stage)
                        state.expected_next_artifact = STAGE_ARTIFACT[state.current_stage].value
                        self.storage.save_state(project_id, state)
                        self.runtime.memory.refresh(state)
                elif action == "compact":
                    with self.storage.transaction(project_id):
                        result = {
                            "context": self.runtime.context_session.compact(project_id),
                            "message": "Working context compacted. Project state, memory and artifacts preserved.",
                        }
                elif action != "status":
                    raise ValueError("Unknown project action")
            finally:
                self.runtime.orchestrator.events.callback = None
            if result.get("error"):
                reply = result["error"]
            elif result.get("questions"):
                reply = "\n".join(result["questions"])
            elif result.get("artifact"):
                a = result["artifact"]
                reply = f"{a.artifact_type} v{a.artifact_version}: {a.approval_status.value}. "
                reply += "Open the document to review, approve or request changes."
            else:
                reply = (
                    result.get("message")
                    or self.storage.load_state(project_id).recommended_next_action
                    or "Ready."
                )
            self._message(
                project_id,
                "assistant",
                reply,
                {
                    "origin": result.get("origin", "workflow"),
                    **({key: result[key] for key in ("provider", "model") if key in result}),
                },
            )
            return {**result, "message": reply, "workspace": self.open_project(project_id)}

    def upload(self, project_id, name, data, artifact_type=None, domain="BUSINESS"):
        from .models import ContextDomain

        domain = ContextDomain(domain).value
        if len(data) > 10 * 1024 * 1024:
            raise ValueError("Maximum upload size is 10 MB")
        suffix = Path(name).suffix.lower()
        if suffix == ".docx":
            from docx import Document

            document = Document(BytesIO(data))
            content = "\n".join(
                [p.text for p in document.paragraphs]
                + [" | ".join(c.text for c in row.cells) for table in document.tables for row in table.rows]
            )
        elif suffix == ".pdf":
            from pypdf import PdfReader

            content = "\n".join(p.extract_text() or "" for p in PdfReader(BytesIO(data)).pages)
        elif suffix in (".txt", ".md", ".json", ".abap", ".csv"):
            content = data.decode("utf-8-sig")
        else:
            content = ""
        with self.lock(project_id), self.storage.transaction(project_id):
            state = self.storage.load_state(project_id)
            file_id = self.storage.save_upload(
                project_id, data, {"name": Path(name).name, "text_extracted": bool(content)}
            )
            if content:
                from .local_storage import DOMAIN_FIELDS

                field = DOMAIN_FIELDS[domain]
                self.storage.save_context(
                    project_id,
                    ContextPackage(
                        project_id=project_id,
                        task_id="upload",
                        **{
                            field: [
                                ContextItem(content=content, source=f"upload:{file_id}", source_type="upload")
                            ]
                        },
                    ),
                )
                if not state.requirement and domain == "BUSINESS":
                    state.requirement = content
            if artifact_type:
                if not content:
                    raise ValueError(
                        "This file needs extracted text before it can be registered as an artifact"
                    )
                kind = ArtifactType(artifact_type)
                versions = [int(r.version) for r in state.artifact_registry if r.artifact_type == kind]
                artifact = Artifact(
                    artifact_id=kind.value,
                    project_id=project_id,
                    artifact_type=kind,
                    content=content,
                    artifact_version=str(max(versions, default=0) + 1),
                    template_version="uploaded",
                    approval_status="draft",
                    edited_by="local",
                    sources=[f"upload:{file_id}"],
                )
                self.storage.save_artifact(project_id, artifact)
                state.artifact_registry.append(artifact.reference())
                artifact_stage = ARTIFACT_STAGE[kind]
                if artifact_stage not in state.started_stages:
                    state.started_stages.append(artifact_stage)
            self.storage.save_state(project_id, state)
            self.runtime.memory.refresh(state)
        return {
            "file_id": file_id,
            "text_extracted": bool(content),
            "workspace": self.open_project(project_id),
        }
