"""Render the implementation decision register as a readable Word document."""

import json
from pathlib import Path
import re

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn

root = Path(__file__).resolve().parents[1]
evidence = json.loads((root / "docs/TEST_RESULTS.json").read_text(encoding="utf-8"))
if evidence["status"] != "passed":
    raise SystemExit("The release gate must pass before the final report is created.")
python_output = next(s["output"] for s in evidence["steps"] if s["name"] == "Backend and CLI full suite")
ui_output = next(s["output"] for s in evidence["steps"] if s["name"] == "UI unit and interaction suite")
python_count = re.search(r"(\d+) passed", python_output).group(1)
ui_count = re.search(r"Tests\s+(\d+) passed", ui_output).group(1)
source = (root / "docs/ASSUMPTIONS_AND_CHANGES.md").read_text(encoding="utf-8")
source = source.replace("{PYTHON_TESTS}", python_count).replace("{UI_TESTS}", ui_count)
(root / "docs/ASSUMPTIONS_AND_CHANGES.md").write_text(source, encoding="utf-8")
document = Document()
for style in document.styles:
    for borders in style.element.findall(".//" + qn("w:pBdr")):
        borders.getparent().remove(borders)
section = document.sections[0]
section.page_width = Inches(8.27)
section.page_height = Inches(11.69)
section.top_margin = section.bottom_margin = Inches(0.72)
section.left_margin = section.right_margin = Inches(0.8)
normal = document.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.08
for name, size in [("Title", 25), ("Heading 1", 19), ("Heading 2", 14)]:
    style = document.styles[name]
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_after = Pt(12)
    style.paragraph_format.space_before = Pt(8)

for chunk_index, chunk in enumerate(source.split("<!-- PAGE -->")):
    if chunk_index:
        document.add_page_break()
    for block in re.split(r"\n\s*\n", chunk.strip()):
        if block.startswith("# "):
            document.add_paragraph(block[2:], "Title")
        elif block.startswith("## "):
            document.add_paragraph(block[3:], "Heading 1")
        else:
            for line in block.splitlines():
                paragraph = document.add_paragraph()
                match = re.match(r"\*\*(.+?)\*\*\s*(.*)", line)
                if match:
                    paragraph.add_run(match.group(1) + " ").bold = True
                    paragraph.add_run(match.group(2))
                else:
                    paragraph.add_run(line)
                paragraph.paragraph_format.widow_control = True

document.core_properties.title = "SAP CoWork V1 Implementation Assumptions and Changes"
document.core_properties.subject = "Implementation decisions and verified acceptance evidence"
document.core_properties.author = "SAP CoWork"
output = root / "docs/SAP_CoWork_V1_Implementation_Assumptions_and_Changes.docx"
document.save(output)
print(output)
